#!/usr/bin/env python3
"""Compare two parts lists in an Excel workbook and write the results.

Usage:
    python compare.py "path/to/workbook.xlsx" [--key MODE]

MODE decides which columns must match for two rows to pair up:
    item-part  Item # + Part # (default; the description is shown but not compared)
    part       Part # only
    full       Item # + Part # + Part Description
    part-desc  Part # + Part Description

The workbook must contain the sheets "List 1" and "List 2", each with the
header row  Item # | Quantity | Part # | Part Description  in row 1 and data
from row 2 down. Results are written as static values to the "Comparison"
sheet (created if missing, cleared first).
"""

import argparse
import sys
from collections import OrderedDict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LIST1 = "List 1"
LIST2 = "List 2"
OUTPUT = "Comparison"

HEADERS = [
    "Item #",
    "Part #",
    "Part Description",
    "Qty List 1",
    "Qty List 2",
    "Difference (L2 - L1)",
    "Status",
]

STATUS_MATCH = "Match"
STATUS_QTY = "Qty Different"
STATUS_ONLY1 = "Only in List 1"
STATUS_ONLY2 = "Only in List 2"
STATUS_ORDER = [STATUS_MATCH, STATUS_QTY, STATUS_ONLY1, STATUS_ONLY2]

FILLS = {
    STATUS_MATCH: PatternFill("solid", fgColor="C6EFCE"),  # green
    STATUS_QTY: PatternFill("solid", fgColor="FFEB9C"),    # yellow
    STATUS_ONLY1: PatternFill("solid", fgColor="FFC7CE"),  # red
    STATUS_ONLY2: PatternFill("solid", fgColor="BDD7EE"),  # blue
}

SUMMARY_COL = 9  # column I; summary block sits beside the table

# key mode -> (use Item #, use Part Description); Part # is always used
KEY_MODES = {
    "full": (True, True),
    "part": (False, False),
    "part-desc": (False, True),
    "item-part": (True, False),
}
DEFAULT_KEY_MODE = "item-part"
KEY_MODE_LABELS = {
    "full": "Item # + Part # + Description",
    "part": "Part # only",
    "part-desc": "Part # + Description",
    "item-part": "Item # + Part #",
}


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------

def clean_text(value):
    """Return the cell value as trimmed text ('' for blank)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_qty(value, where):
    """Coerce a quantity cell to a number. Blank or invalid -> 0."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else value
    try:
        num = float(str(value).strip().replace(",", ""))
    except ValueError:
        print(f"  warning: {where}: quantity {value!r} is not a number, using 0")
        return 0
    return int(num) if num.is_integer() else num


def read_list(ws, key_mode=DEFAULT_KEY_MODE):
    """Read a list sheet into rows: (key, display_item, display_part,
    display_desc, qty). Fully blank rows are skipped."""
    use_item, use_desc = KEY_MODES[key_mode]
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
        item, qty, part, desc = (list(row) + [None] * 4)[:4]
        item_t, part_t, desc_t = clean_text(item), clean_text(part), clean_text(desc)
        if not item_t and not part_t and not desc_t and qty is None:
            continue
        key = (item_t.lower() if use_item else "",
               part_t.lower(),
               desc_t.lower() if use_desc else "")
        rows.append((key, item_t, part_t, desc_t,
                     parse_qty(qty, f"{ws.title} row {idx}")))
    return rows


def group_by_key(rows):
    """Map key -> list of (display_item, display_part, display_desc, qty),
    preserving order of appearance so duplicates pair 1st-with-1st, etc."""
    groups = OrderedDict()
    for key, item, part, desc, qty in rows:
        groups.setdefault(key, []).append((item, part, desc, qty))
    return groups


# --------------------------------------------------------------------------
# Comparison
# --------------------------------------------------------------------------

def sort_token(text):
    """Sort numerically when the text is a number, otherwise as text."""
    try:
        return (0, float(text), "")
    except ValueError:
        return (1, 0.0, text.lower())


def compare(list1_rows, list2_rows):
    g1 = group_by_key(list1_rows)
    g2 = group_by_key(list2_rows)

    results = []
    for key in list(g1.keys()) + [k for k in g2 if k not in g1]:
        occ1 = g1.get(key, [])
        occ2 = g2.get(key, [])
        for i in range(max(len(occ1), len(occ2))):
            a = occ1[i] if i < len(occ1) else None
            b = occ2[i] if i < len(occ2) else None
            src = a or b
            item, part, desc = src[0], src[1], src[2]
            q1 = a[3] if a else 0
            q2 = b[3] if b else 0
            if a and b:
                status = STATUS_MATCH if q1 == q2 else STATUS_QTY
            elif a:
                status = STATUS_ONLY1
            else:
                status = STATUS_ONLY2
            results.append({
                "item": item, "part": part, "desc": desc,
                "q1": q1, "q2": q2, "diff": q2 - q1, "status": status,
                "_key": key, "_occ": i,
            })

    results.sort(key=lambda r: (sort_token(r["item"]), sort_token(r["part"]),
                                r["desc"].lower(), r["_occ"]))
    return results


# --------------------------------------------------------------------------
# Writing
# --------------------------------------------------------------------------

def clear_sheet(ws):
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws.auto_filter.ref = None
    ws.freeze_panes = None
    ws.conditional_formatting = type(ws.conditional_formatting)()
    ws.merged_cells.ranges = set()
    for col in list(ws.column_dimensions.keys()):
        del ws.column_dimensions[col]


def write_results(ws, results, key_mode=DEFAULT_KEY_MODE):
    clear_sheet(ws)
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9D9D9")

    # Header row
    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for r, row in enumerate(results, start=2):
        values = [row["item"], row["part"], row["desc"],
                  row["q1"], row["q2"], row["diff"], row["status"]]
        fill = FILLS[row["status"]]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.fill = fill

    last_row = max(len(results) + 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last_row}"
    ws.freeze_panes = "A2"

    # Summary block beside the table
    counts = {s: 0 for s in STATUS_ORDER}
    for row in results:
        counts[row["status"]] += 1

    ws.cell(row=1, column=SUMMARY_COL, value="Summary").font = bold
    ws.cell(row=1, column=SUMMARY_COL + 1, value="Count").font = bold
    for i, status in enumerate(STATUS_ORDER, start=2):
        label = ws.cell(row=i, column=SUMMARY_COL, value=status)
        label.fill = FILLS[status]
        ws.cell(row=i, column=SUMMARY_COL + 1, value=counts[status]).fill = FILLS[status]
    total_row = len(STATUS_ORDER) + 2
    ws.cell(row=total_row, column=SUMMARY_COL, value="Total").font = bold
    ws.cell(row=total_row, column=SUMMARY_COL + 1, value=len(results)).font = bold
    ws.cell(row=total_row + 2, column=SUMMARY_COL, value="Matched on").font = bold
    ws.cell(row=total_row + 2, column=SUMMARY_COL + 1, value=KEY_MODE_LABELS[key_mode])

    autofit(ws)
    return counts


def autofit(ws, min_width=8, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(
            min_width, min(width + 2, max_width))


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def run(path, key_mode=DEFAULT_KEY_MODE):
    wb = load_workbook(path)
    for name in (LIST1, LIST2):
        if name not in wb.sheetnames:
            sys.exit(f"error: sheet '{name}' not found in {path}")

    list1 = read_list(wb[LIST1], key_mode)
    list2 = read_list(wb[LIST2], key_mode)
    results = compare(list1, list2)

    ws = wb[OUTPUT] if OUTPUT in wb.sheetnames else wb.create_sheet(OUTPUT)
    counts = write_results(ws, results, key_mode)
    wb.save(path)

    print(f"Read {len(list1)} rows from '{LIST1}', {len(list2)} rows from '{LIST2}'.")
    print(f"Matched on: {KEY_MODE_LABELS[key_mode]}")
    print(f"Wrote {len(results)} rows to '{OUTPUT}' in {path}")
    for status in STATUS_ORDER:
        print(f"  {status:<15} {counts[status]}")
    return results


def main(argv):
    ap = argparse.ArgumentParser(description="Compare List 1 and List 2 in a workbook.")
    ap.add_argument("workbook", help="path to the .xlsx workbook")
    ap.add_argument("--key", choices=sorted(KEY_MODES), default=DEFAULT_KEY_MODE,
                    help="columns that must match for rows to pair up "
                         "(item-part = Item # + Part #, the default; part = Part # only; "
                         "full = Item # + Part # + Description; part-desc = Part # + Description)")
    args = ap.parse_args(argv[1:])
    run(args.workbook, args.key)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
