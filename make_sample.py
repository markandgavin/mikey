#!/usr/bin/env python3
"""Create the blank template and the sample workbook used to test compare.py.

Usage:
    python make_sample.py

Writes:
    parts_comparison_template.xlsx  - empty List 1 / List 2 / Comparison tabs
    sample_parts.xlsx               - test data covering every Status case
"""

from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = ["Item #", "Quantity", "Part #", "Part Description"]

# Expected result per row is noted in the comment.
LIST1 = [
    ("100", 4, "PN-001", "Bolt M6"),             # Match
    ("101", 2, "PN-002", "Nut M6"),              # Qty Different (2 vs 5)
    ("102", 1, "PN-003", "Washer M6"),           # Only in List 1
    ("103", 3, "pn-004", "Bracket Left  "),      # Match (case / whitespace differ)
    ("104", 2, "PN-005", "Gasket"),              # dup #1 -> Match
    ("104", 2, "PN-005", "Gasket"),              # dup #2 -> Only in List 1
    ("105", 6, "PN-006", "Hinge"),               # Match (List 2 has a 2nd copy)
    ("106", 1, "PN-007", "Screw M4"),            # Match
    ("107", 10, "PN-008", "Cable 2m"),           # Qty Different (10 vs 8)
    ("108", 1, "PN-009", "Fuse 5A"),             # Only in List 1
    ("109", None, "PN-010", "Relay 12V"),        # Qty Different (blank -> 0 vs 2)
    ("112", 1, "PN-013", "Pump A"),              # Only in List 1 (desc differs)
]

LIST2 = [
    ("100", 4, "PN-001", "Bolt M6"),
    ("101", 5, "PN-002", "Nut M6"),
    ("103", 3, "PN-004", "bracket left"),
    ("104", 2, "PN-005", "Gasket"),
    ("105", 6, "PN-006", "Hinge"),
    ("105", 7, "PN-006", "Hinge"),               # dup #2 -> Only in List 2
    ("106", 1, "PN-007", "Screw M4"),
    ("107", 8, "PN-008", "Cable 2m"),
    ("109", 2, "PN-010", "Relay 12V"),
    ("110", 3, "PN-011", "Sensor"),              # Only in List 2
    ("112", 1, "PN-013", "Pump B"),              # Only in List 2 (desc differs)
]


def build(list1, list2):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "List 1"
    ws2 = wb.create_sheet("List 2")
    wb.create_sheet("Comparison")
    for ws, rows in ((ws1, list1), (ws2, list2)):
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))
        for col, width in zip("ABCD", (10, 10, 12, 30)):
            ws.column_dimensions[col].width = width
    return wb


if __name__ == "__main__":
    build([], []).save("parts_comparison_template.xlsx")
    build(LIST1, LIST2).save("sample_parts.xlsx")
    print("Wrote parts_comparison_template.xlsx and sample_parts.xlsx")
