#!/usr/bin/env python3
"""Regression test for compare.py against the sample workbook.

Usage:
    python test_compare.py
"""

import os
import shutil
import tempfile

from openpyxl import load_workbook

import compare
import make_sample

EXPECTED = [
    # item, part, desc, q1, q2, diff, status
    ("100", "PN-001", "Bolt M6", 4, 4, 0, "Match"),
    ("101", "PN-002", "Nut M6", 2, 5, 3, "Qty Different"),
    ("102", "PN-003", "Washer M6", 1, 0, -1, "Only in List 1"),
    ("103", "pn-004", "Bracket Left", 3, 3, 0, "Match"),
    ("104", "PN-005", "Gasket", 2, 2, 0, "Match"),
    ("104", "PN-005", "Gasket", 2, 0, -2, "Only in List 1"),
    ("105", "PN-006", "Hinge", 6, 6, 0, "Match"),
    ("105", "PN-006", "Hinge", 0, 7, 7, "Only in List 2"),
    ("106", "PN-007", "Screw M4", 1, 1, 0, "Match"),
    ("107", "PN-008", "Cable 2m", 10, 8, -2, "Qty Different"),
    ("108", "PN-009", "Fuse 5A", 1, 0, -1, "Only in List 1"),
    ("109", "PN-010", "Relay 12V", 0, 2, 2, "Qty Different"),
    ("110", "PN-011", "Sensor", 0, 3, 3, "Only in List 2"),
    ("112", "PN-013", "Pump A", 1, 1, 0, "Match"),
]

EXPECTED_COUNTS = {"Match": 6, "Qty Different": 3,
                   "Only in List 1": 3, "Only in List 2": 2}


def main():
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, "sample.xlsx")
    make_sample.build(make_sample.LIST1, make_sample.LIST2).save(path)

    # Run twice: the second run must clear the first run's output.
    compare.run(path)
    compare.run(path)

    wb = load_workbook(path)
    ws = wb["Comparison"]

    header = [c.value for c in ws[1]][:7]
    assert header == compare.HEADERS, header

    rows = [tuple(c.value for c in r[:7])
            for r in ws.iter_rows(min_row=2, max_row=ws.max_row)
            if r[0].value is not None]
    assert rows == EXPECTED, "\n".join(f"{a}\n{b}" for a, b in zip(rows, EXPECTED) if a != b)
    assert ws.max_row == len(EXPECTED) + 1, ws.max_row

    # Formatting
    assert all(c.font.bold for c in ws[1][:7])
    assert ws.freeze_panes == "A2", ws.freeze_panes
    assert ws.auto_filter.ref == f"A1:G{len(EXPECTED) + 1}", ws.auto_filter.ref
    for r, exp in enumerate(EXPECTED, start=2):
        colour = ws.cell(row=r, column=1).fill.fgColor.rgb
        assert colour == compare.FILLS[exp[6]].fgColor.rgb, (r, colour)
    assert all(ws.column_dimensions[col].width >= 8 for col in "ABCDEFG")

    # Summary block (columns I:J)
    summary = {ws.cell(row=r, column=9).value: ws.cell(row=r, column=10).value
               for r in range(2, 7)}
    assert summary == {**EXPECTED_COUNTS, "Total": len(EXPECTED)}, summary
    assert ws.cell(row=8, column=10).value == "Item # + Part #"

    # Full key (description compared) splits the Pump A / Pump B pair
    make_sample.build(make_sample.LIST1, make_sample.LIST2).save(path)
    compare.run(path, "full")
    full = load_workbook(path)["Comparison"]
    statuses = [r[6] for r in full.iter_rows(min_row=2, max_row=full.max_row, max_col=7, values_only=True)]
    assert statuses.count("Only in List 1") == 4 and statuses.count("Only in List 2") == 3, statuses

    # No formulas anywhere
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("=")), c.coordinate

    shutil.rmtree(tmp)
    print(f"OK: {len(rows)} rows, counts {EXPECTED_COUNTS}")


if __name__ == "__main__":
    main()
