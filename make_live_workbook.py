#!/usr/bin/env python3
"""Build the formula-driven ("live") comparison workbook.

Unlike compare.py, this workbook needs no script: paste data into List 1
and List 2 and the Comparison tab recalculates by itself. It uses only
classic Excel functions (no dynamic arrays, no macros), so it works in
Excel desktop, Excel for the web, and the Excel phone apps.

Usage:
    python make_live_workbook.py            # blank + sample workbooks
    python make_live_workbook.py --rows 800 # allow up to 800 rows per list

Writes:
    parts_comparison_live.xlsx   - blank, ready for pasting
    sample_parts_live.xlsx       - sample data preloaded
"""

import argparse
import os
import re
import shutil
import subprocess
import tempfile
import zipfile

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import make_sample
from compare import FILLS, HEADERS, STATUS_ORDER

LIST_HEADERS = ["Item #", "Quantity", "Part #", "Part Description"]


def build(list1_rows, list2_rows, n):
    """n = maximum data rows per list. The Comparison tab holds 2n rows."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "List 1"
    ws2 = wb.create_sheet("List 2")
    cmp_ws = wb.create_sheet("Comparison")
    calc = wb.create_sheet("Calc")

    bold = Font(bold=True)
    for ws, rows in ((ws1, list1_rows), (ws2, list2_rows)):
        ws.append(LIST_HEADERS)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(list(row))
        for col, width in zip("ABCD", (10, 10, 12, 30)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"

    last = n + 1          # last source row in each list
    m = 2 * n + 1         # last union row

    # ------------------------------------------------------------------
    # Calc sheet (hidden). Columns:
    #   A key1  B occ1  C keyocc1  D qty1          (List 1, rows 2..n+1)
    #   F key2  G occ2  H keyocc2  I qty2          (List 2, rows 2..n+1)
    #   K keyocc  L item  M part  N desc  O q1  P q2  Q in1  R in2
    #   S status  T itemKey  U partKey  V descKey  W occ  X rank
    #                                              (union, rows 2..2n+1)
    # ------------------------------------------------------------------
    calc.append(["key1", "occ1", "keyocc1", "qty1", "",
                 "key2", "occ2", "keyocc2", "qty2", "",
                 "keyocc", "Item #", "Part #", "Part Description", "q1", "q2",
                 "in1", "in2", "Status", "itemKey", "partKey", "descKey",
                 "occ", "rank"])

    def key_block(sheet, kcol, ocol, kocol, qcol):
        for i in range(2, last + 1):
            s = f"'{sheet}'!"
            calc[f"{kcol}{i}"] = (
                f'=IF(TRIM({s}A{i}&{s}C{i}&{s}D{i})="","",'
                f'LOWER(TRIM({s}A{i})&"|"&TRIM({s}C{i})&"|"&TRIM({s}D{i})))'
            )
            calc[f"{ocol}{i}"] = (
                f'=IF({kcol}{i}="","",SUMPRODUCT(--(${kcol}$2:{kcol}{i}={kcol}{i})))'
            )
            calc[f"{kocol}{i}"] = f'=IF({kcol}{i}="","",{kcol}{i}&"#"&{ocol}{i})'
            calc[f"{qcol}{i}"] = f'=IF({kcol}{i}="","",IFERROR(--{s}B{i},0))'

    key_block("List 1", "A", "B", "C", "D")
    key_block("List 2", "F", "G", "H", "I")

    def display(src_sheet, src_col, j, i):
        s = f"'{src_sheet}'!{src_col}{j}"
        return f'=IF(K{i}="","",IF(ISNUMBER({s}),{s},TRIM({s})))'

    for i in range(2, m + 1):
        from_list1 = i <= last
        j = i if from_list1 else i - n
        sheet = "List 1" if from_list1 else "List 2"

        if from_list1:
            calc[f"K{i}"] = f"=C{j}"
            calc[f"O{i}"] = f'=IF(K{i}="","",D{j})'
            calc[f"P{i}"] = (f'=IF(K{i}="","",IF(R{i},'
                             f'SUMPRODUCT(--($H$2:$H${last}=K{i}),$I$2:$I${last}),0))')
            calc[f"Q{i}"] = f'=IF(K{i}="","",TRUE)'
            calc[f"R{i}"] = f'=IF(K{i}="","",SUMPRODUCT(--($H$2:$H${last}=K{i}))>0)'
            calc[f"W{i}"] = f'=IF(K{i}="","",B{j})'
        else:
            calc[f"K{i}"] = (f'=IF(H{j}="","",'
                             f'IF(SUMPRODUCT(--($C$2:$C${last}=H{j}))>0,"",H{j}))')
            calc[f"O{i}"] = f'=IF(K{i}="","",0)'
            calc[f"P{i}"] = f'=IF(K{i}="","",I{j})'
            calc[f"Q{i}"] = f'=IF(K{i}="","",FALSE)'
            calc[f"R{i}"] = f'=IF(K{i}="","",TRUE)'
            calc[f"W{i}"] = f'=IF(K{i}="","",G{j})'

        calc[f"L{i}"] = display(sheet, "A", j, i)
        calc[f"M{i}"] = display(sheet, "C", j, i)
        calc[f"N{i}"] = display(sheet, "D", j, i)
        calc[f"S{i}"] = (
            f'=IF(K{i}="","",IF(AND(Q{i},R{i}),'
            f'IF(O{i}=P{i},"{STATUS_ORDER[0]}","{STATUS_ORDER[1]}"),'
            f'IF(Q{i},"{STATUS_ORDER[2]}","{STATUS_ORDER[3]}")))'
        )
        calc[f"T{i}"] = (f'=IF(K{i}="","",IF(ISNUMBER(--L{i}),'
                         f'TEXT(--L{i},"000000000000.000"),LOWER(L{i})))')
        calc[f"U{i}"] = (f'=IF(K{i}="","",IF(ISNUMBER(--M{i}),'
                         f'TEXT(--M{i},"000000000000.000"),LOWER(M{i})))')
        calc[f"V{i}"] = f'=IF(K{i}="","",LOWER(N{i}))'
        rng = lambda c: f"${c}$2:${c}${m}"  # noqa: E731
        calc[f"X{i}"] = (
            f'=IF(K{i}="","",1+SUMPRODUCT(({rng("K")}<>"")*('
            f'({rng("T")}<T{i})+({rng("T")}=T{i})*('
            f'({rng("U")}<U{i})+({rng("U")}=U{i})*('
            f'({rng("V")}<V{i})+({rng("V")}=V{i})*('
            f'({rng("W")}<W{i})+({rng("W")}=W{i})*(ROW({rng("K")})<ROW(K{i}))'
            f'))))))'
        )
    calc.sheet_state = "hidden"

    # ------------------------------------------------------------------
    # Comparison sheet
    # ------------------------------------------------------------------
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    for c, name in enumerate(HEADERS, start=1):
        cell = cmp_ws.cell(row=1, column=c, value=name)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    src = {"A": "L", "B": "M", "C": "N", "D": "O", "E": "P", "G": "S"}
    for r in range(2, m + 1):
        for out_col, calc_col in src.items():
            cmp_ws[f"{out_col}{r}"] = (
                f'=IFERROR(INDEX(Calc!${calc_col}$2:${calc_col}${m},'
                f'MATCH(ROW()-1,Calc!$X$2:$X${m},0)),"")'
            )
        cmp_ws[f"F{r}"] = f'=IF(G{r}="","",E{r}-D{r})'

    cmp_ws.auto_filter.ref = f"A1:G{m}"
    cmp_ws.freeze_panes = "A2"
    for status in STATUS_ORDER:
        cmp_ws.conditional_formatting.add(
            f"A2:G{m}",
            FormulaRule(formula=[f'$G2="{status}"'], fill=FILLS[status]),
        )

    # Summary block beside the table (I:J)
    cmp_ws["I1"] = "Summary"
    cmp_ws["J1"] = "Count"
    cmp_ws["I1"].font = bold
    cmp_ws["J1"].font = bold
    for i, status in enumerate(STATUS_ORDER, start=2):
        cmp_ws[f"I{i}"] = status
        cmp_ws[f"J{i}"] = f'=COUNTIF($G$2:$G${m},I{i})'
        cmp_ws[f"I{i}"].fill = FILLS[status]
        cmp_ws[f"J{i}"].fill = FILLS[status]
    total_row = len(STATUS_ORDER) + 2
    cmp_ws[f"I{total_row}"] = "Total"
    cmp_ws[f"J{total_row}"] = f"=SUM(J2:J{total_row - 1})"
    cmp_ws[f"I{total_row}"].font = bold
    cmp_ws[f"J{total_row}"].font = bold

    for col, width in zip("ABCDEFGHIJ", (10, 12, 30, 11, 11, 20, 16, 3, 16, 8)):
        cmp_ws.column_dimensions[col].width = width
    for r in range(2, m + 1):
        for col in "DEF":
            cmp_ws[f"{col}{r}"].alignment = Alignment(horizontal="right")

    return wb


def bake_values(path):
    """Recalculate the workbook with LibreOffice so every formula cell carries
    a cached result. Without this, viewers that do not recalculate on open
    (file previews, some phone apps) show the Comparison tab empty.
    Returns True when the values were baked in, False if LibreOffice is absent."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        env = dict(os.environ, HOME=tmp)  # LibreOffice needs a writable profile
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx",
                        "--outdir", tmp, os.path.abspath(path)],
                       check=True, env=env, capture_output=True, timeout=600)
        out = os.path.join(tmp, os.path.basename(path))
        if not os.path.exists(out):
            return False
        # Keep "recalculate on load" so Excel refreshes results after any edit
        # even if a viewer only reads the cached ones.
        patched = path + ".tmp"
        with zipfile.ZipFile(out) as src, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "xl/workbook.xml":
                    text = data.decode("utf-8")
                    if "fullCalcOnLoad" not in text:
                        text = re.sub(r"<calcPr", '<calcPr fullCalcOnLoad="1"', text, count=1)
                    data = text.encode("utf-8")
                dst.writestr(item, data)
        os.replace(patched, path)
    return True


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--rows", type=int, default=500,
                    help="maximum rows per list (default 500)")
    ap.add_argument("--no-bake", action="store_true",
                    help="skip the LibreOffice recalculation step")
    args = ap.parse_args()
    outputs = {
        "parts_comparison_live.xlsx": ([], []),
        "sample_parts_live.xlsx": (make_sample.LIST1, make_sample.LIST2),
    }
    for name, (l1, l2) in outputs.items():
        build(l1, l2, args.rows).save(name)
        baked = False if args.no_bake else bake_values(name)
        print(f"Wrote {name} ({args.rows} rows per list, "
              f"{'results baked in' if baked else 'no cached results: LibreOffice not found'})")


if __name__ == "__main__":
    main()
